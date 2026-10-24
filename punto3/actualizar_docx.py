#!/usr/bin/env python3
"""
Actualiza el documento Word oficial:
'C:\\Users\\samue\\OneDrive\\Documentos\\Obsidian Vault\\Taller\\Taller 3 - SAMUEL ZULUAGA SAMUEL ORTIZ JUAN MEJIA.docx'
insertando el Desarrollo del Punto 3 y la sección de IA para el Punto 3.
"""

import os
import sys
import docx
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement, parse_xml
from docx.oxml.ns import nsdecls, qn

DOC_PATH = r"C:\Users\samue\OneDrive\Documentos\Obsidian Vault\Taller\Taller 3 - SAMUEL ZULUAGA SAMUEL ORTIZ JUAN MEJIA.docx"
GRAFICA1_PATH = r"C:\Users\samue\Downloads\TallerOyB\punto3\graficas\1_comparativos_vs_no_comparativos.png"
GRAFICA2_PATH = r"C:\Users\samue\Downloads\TallerOyB\punto3\graficas\2_ordenar_vs_no_ordenar_breakeven.png"

def set_cell_shading(cell, color_hex="E7E6E6"):
    shading_xml = f'<w:shd {nsdecls("w")} w:fill="{color_hex}" w:val="clear"/>'
    cell._tc.get_or_add_tcPr().append(parse_xml(shading_xml))

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    tcPr = cell._tc.get_or_add_tcPr()
    tcMar = OxmlElement('w:tcMar')
    for m, val in [('top', top), ('bottom', bottom), ('left', left), ('right', right)]:
        node = OxmlElement(f'w:{m}')
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tcMar.append(node)
    tcPr.append(tcMar)

def style_table(table, col_widths, col_alignments):
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    # tblPr borders
    tblPr = table._tbl.tblPr
    borders_xml = f'''
    <w:tblBorders {nsdecls("w")}>
        <w:top w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:left w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:bottom w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:right w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:insideH w:val="single" w:sz="4" w:space="0" w:color="000000"/>
        <w:insideV w:val="single" w:sz="4" w:space="0" w:color="000000"/>
    </w:tblBorders>
    '''
    tblPr.append(parse_xml(borders_xml))
    
    for row_idx, row in enumerate(table.rows):
        is_header = (row_idx == 0)
        # Prevent row split across pages
        trPr = row._tr.get_or_add_trPr()
        trPr.append(parse_xml(f'<w:cantSplit {nsdecls("w")}/>'))
        if is_header:
            trPr.append(parse_xml(f'<w:tblHeader {nsdecls("w")}/>'))
            
        for col_idx, cell in enumerate(row.cells):
            cell.width = col_widths[col_idx]
            cell.vertical_alignment = WD_ALIGN_VERTICAL.CENTER
            set_cell_margins(cell, top=120, bottom=120, left=140, right=140)
            if is_header:
                set_cell_shading(cell, "E7E6E6")
            for p in cell.paragraphs:
                p.alignment = col_alignments[col_idx]
                p.paragraph_format.space_before = Pt(2)
                p.paragraph_format.space_after = Pt(2)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.name = "Calibri"
                    run.font.size = Pt(8.5) if len(col_widths) > 5 else Pt(9.0)
                    if is_header:
                        run.font.bold = True

def add_p(ref_p, text="", bold=False, italic=False, size_pt=10.0, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=3, space_after=3, color_rgb=None, heading_style=None):
    new_p = ref_p.insert_paragraph_before()
    new_p.alignment = align
    new_p.paragraph_format.space_before = Pt(space_before)
    new_p.paragraph_format.space_after = Pt(space_after)
    new_p.paragraph_format.line_spacing = 1.15
    if heading_style:
        new_p.style = heading_style
    if text:
        run = new_p.add_run(text)
        run.font.name = "Calibri"
        run.font.size = Pt(size_pt)
        run.font.bold = bold
        run.font.italic = italic
        if color_rgb:
            run.font.color.rgb = color_rgb
    return new_p

def add_custom_p(ref_p, runs_data, align=WD_ALIGN_PARAGRAPH.JUSTIFY, space_before=3, space_after=3):
    new_p = ref_p.insert_paragraph_before()
    new_p.alignment = align
    new_p.paragraph_format.space_before = Pt(space_before)
    new_p.paragraph_format.space_after = Pt(space_after)
    new_p.paragraph_format.line_spacing = 1.15
    for r_text, r_bold, r_italic, r_size, r_color in runs_data:
        run = new_p.add_run(r_text)
        run.font.name = "Calibri"
        run.font.size = Pt(r_size)
        run.font.bold = r_bold
        run.font.italic = r_italic
        if r_color:
            run.font.color.rgb = r_color
    return new_p

def ejecutar():
    print(f"Cargando documento desde: {DOC_PATH}")
    doc = docx.Document(DOC_PATH)
    
    # Sanitizar pgMar en caso de atributos con decimales (exportados por herramientas externas)
    for mar in doc._element.xpath('//w:pgMar'):
        for k, v in list(mar.attrib.items()):
            if '.' in v:
                mar.attrib[k] = str(round(float(v)))
    
    # 1. Localizar el punto de inserción para el Desarrollo del Punto 3 (antes de 'Entrega')
    target_idx = None
    for i, p in enumerate(doc.paragraphs):
        if p.text.strip().startswith("Entrega"):
            target_idx = i
            break
            
    if target_idx is None:
        raise ValueError("No se encontró el párrafo 'Entrega' en el documento.")
        
    print(f"Punto de inserción localizado en párrafo {target_idx}: '{doc.paragraphs[target_idx].text[:30]}'")
    ref_p = doc.paragraphs[target_idx]
    
    # -------------------------------------------------------------
    # SECCIÓN: DESARROLLO DEL PUNTO 3: ANÁLISIS COMPARATIVO
    # -------------------------------------------------------------
    add_p(ref_p, "Desarrollo del Punto 3: Análisis Comparativo", bold=True, size_pt=11.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=14, space_after=8)
    
    # --- SUBSECCIÓN 1 ---
    add_p(ref_p, "1. Comparativos frente a No Comparativos", bold=True, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=10, space_after=4, heading_style="Heading 1")
    
    add_p(ref_p, "En esta sección se investiga el fundamento teórico que separa a los algoritmos basados en comparaciones de aquellos que operan sin comparar elementos, contrastando los postulados formales con los tiempos reales medidos en el Punto 1 para n = 1.000.000 de elementos.")
    
    # A. Límite inferior teórico
    add_p(ref_p, "Límite inferior teórico para algoritmos basados en comparaciones (Cota Ω(n log n)):", bold=True, size_pt=10.0, space_before=6, space_after=2)
    
    add_p(ref_p, "En el modelo formal de árbol de decisión (Decision Tree Model), cualquier algoritmo de ordenamiento que opere exclusivamente mediante comparaciones binarias (del tipo a_i ≤ a_j) puede modelarse como un árbol binario donde cada nodo interno representa una comparación y cada hoja representa una de las posibles permutaciones de los datos de entrada.")
    
    add_custom_p(ref_p, [
        ("1. Cantidad de permutaciones: ", True, False, 10.0, None),
        ("Un arreglo de n elementos distintos posee exactamente n! (factorial de n) ordenamientos o permutaciones posibles. Para garantizar que el algoritmo ordene correctamente cualquier permutación arbitraria, el árbol de decisión debe poseer al menos tantas hojas (L) como permutaciones existan, es decir, L ≥ n!.", False, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("2. Altura del árbol y número de comparaciones: ", True, False, 10.0, None),
        ("En un árbol binario, un árbol de altura h puede albergar como máximo 2^h hojas. Por consiguiente, para satisfacer que 2^h ≥ L ≥ n!, la altura mínima del árbol debe ser h ≥ log₂(n!).", False, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("3. Deducción asintótica por la aproximación de Stirling: ", True, False, 10.0, None),
        ("Aplicando la cota inferior factorial, sabemos que n! ≥ (n/2)^(n/2). Al tomar logaritmo en base 2: log₂(n!) ≥ (n/2) · log₂(n/2) = (n/2) · (log₂ n - 1) = Ω(n log n). De forma más exacta por la aproximación de Stirling, ln(n!) = n ln n - n + O(ln n), lo que conduce a que log₂(n!) = n log₂ n - 1,4427n + O(log n).", False, False, 10.0, None)
    ])
    
    add_p(ref_p, "Conclusión: La altura del árbol de decisión representa el número de comparaciones en el peor caso. Ningún algoritmo de ordenamiento basado en comparaciones puede tener una complejidad temporal en el peor caso inferior a Ω(n log n). Algoritmos óptimos en este modelo como Merge Sort y Heap Sort alcanzan exactamente esta cota inferior.")
    
    # B. Por qué Radix y Bucket escapan a este límite
    add_p(ref_p, "¿Por qué Radix Sort y Bucket Sort no están sujetos a la cota Ω(n log n)?:", bold=True, size_pt=10.0, space_before=6, space_after=2)
    add_p(ref_p, "Radix Sort y Bucket Sort logran superar la barrera teórica de Ω(n log n) y alcanzar una complejidad en tiempo lineal O(n) porque no operan mediante comparaciones relativas entre elementos. En su lugar, explotan la estructura interna, la representación posicional de los datos y el direccionamiento indexado en memoria:")
    
    add_custom_p(ref_p, [
        ("• Radix Sort (LSD): ", True, False, 10.0, None),
        ("Trata cada clave como una secuencia de d dígitos en una base b. Aplica de forma sucesiva un ordenamiento auxiliar estable (Counting Sort) sobre cada dígito posicional, desde el menos significativo (LSD) hasta el más significativo. Counting Sort realiza conteos de frecuencias e indexación aritmética directa en memoria en tiempo O(n + b). Como no compara elementos entre sí, para d dígitos fijos el costo total es O(d · (n + b)) = O(n).", False, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("• Bucket Sort: ", True, False, 10.0, None),
        ("Asume que los datos pertenecen a un intervalo acotado [min, max] y divide el rango en k cubetas disjuntas. Mapea cada elemento a su cubeta en tiempo O(1) mediante una función de dispersión aritmética directa (índice = ⌊((arr[i] - min) / rango) · k⌋). Al ordenar luego cada cubeta y concatenarlas, si los elementos se distribuyen de manera uniforme, cada cubeta contendrá en promedio λ = n/k elementos. Cuando k = Θ(n), el tiempo esperado de ordenamiento total es estrictamente lineal O(n).", False, False, 10.0, None)
    ])
    
    # C. Comparación cuantitativa con mediciones
    add_p(ref_p, "Comparación de los tiempos experimentales medidos (n = 1.000.000, entrada aleatoria):", bold=True, size_pt=10.0, space_before=6, space_after=2)
    add_p(ref_p, "La Tabla 17 y la Figura 10 comparan los tiempos obtenidos en el Punto 1 para ordenar 1.000.000 de enteros de 8 dígitos generados pseudoaleatoriamente.")
    
    add_p(ref_p, "Tabla 17. Comparación de tiempos entre familias de ordenamiento (n = 1.000.000, aleatorio)", bold=True, size_pt=9.0, space_before=4, space_after=2)
    
    # Tabla 17
    # Columnas: Algoritmo, Familia, Tipo, Complejidad Teórica, Java (ms), Python (ms), Python / Java
    tbl17 = doc.add_table(rows=7, cols=7)
    ref_p._p.addprevious(tbl17._tbl)
    
    headers17 = ["Algoritmo", "Familia", "Tipo", "Complejidad Teórica", "Java (ms)", "Python (ms)", "Python / Java"]
    filas17 = [
        ["Radix Sort", "Distribución", "No comparativo", "O(d · (n + b)) ~ O(n)", "83,320 ± 9,307", "5.951,966 ± 117,606", "71,4×"],
        ["Bucket Sort", "Distribución", "No comparativo", "O(n) prom / O(n²) peor", "498,618 ± 128,993", "1.915,109 ± 95,585", "3,8×"],
        ["Dual-Pivot QuickSort", "Partición", "Comparativo", "O(n log n) / O(n²) peor", "91,854 ± 7,004", "2.600,673 ± 80,337", "28,3×"],
        ["Heap Sort", "Selección", "Comparativo", "O(n log n)", "175,949 ± 12,364", "8.585,681 ± 503,718", "48,8×"],
        ["Merge Sort", "Mezcla", "Comparativo", "O(n log n)", "192,903 ± 24,911", "6.090,020 ± 305,685", "31,6×"],
        ["Insertion Sort", "Inserción", "Comparativo", "O(n²)", "78.826,821 ± 1.429,596", "30.575.356,2 (est.)", "388× (est.)"]
    ]
    
    for c_idx, h_text in enumerate(headers17):
        tbl17.cell(0, c_idx).text = h_text
    for r_idx, fila in enumerate(filas17):
        for c_idx, val in enumerate(fila):
            tbl17.cell(r_idx + 1, c_idx).text = val
            
    col_w17 = [Inches(1.25), Inches(0.9), Inches(1.05), Inches(1.3), Inches(1.1), Inches(1.1), Inches(0.7)]
    col_a17 = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
    style_table(tbl17, col_w17, col_a17)
    
    # Figura 10
    p_img10 = ref_p.insert_paragraph_before()
    p_img10.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img10.paragraph_format.space_before = Pt(6)
    p_img10.paragraph_format.space_after = Pt(2)
    run10 = p_img10.add_run()
    run10.add_picture(GRAFICA1_PATH, width=Inches(6.2))
    
    add_p(ref_p, "Figura 10. Algoritmos de ordenamiento para n = 1.000.000: Comparativos frente a No Comparativos en Java y Python.", italic=True, size_pt=9.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=6)
    
    add_p(ref_p, "Análisis de los resultados empíricos:")
    add_custom_p(ref_p, [
        ("1. Radix Sort lideró en Java (83,32 ms): ", True, False, 10.0, None),
        ("Superó con solvencia al algoritmo comparativo más rápido (Dual-Pivot QuickSort con 91,85 ms), así como a Heap Sort (175,95 ms) y Merge Sort (192,90 ms). En Java, el compilador JIT traduce las 8 pasadas de Counting Sort a código nativo sobre arreglos primitivos int[] de memoria contigua sin sobrecarga de objetos ni penalización de predicción de saltos condicionales.", False, False, 10.0, None)
    ])
    add_custom_p(ref_p, [
        ("2. Bucket Sort lideró en Python (1.915,11 ms): ", True, False, 10.0, None),
        ("Fue el algoritmo más rápido de todo el benchmark en Python, aventajando a Dual-Pivot QuickSort (2.600,67 ms) y a Radix Sort (5.951,97 ms). En Python, la interpretación de bucles es muy costosa; Radix Sort sufre porque debe ejecutar 8 recorridos interpretados completos sobre un millón de elementos, mientras que Bucket Sort realiza una única pasada de dispersión y reparte los datos en 100.000 cubetas con un promedio de sólo 10 elementos por cubeta, donde Insertion Sort corre casi instantáneamente.", False, False, 10.0, None)
    ])
    
    # D. Supuestos sobre los datos de entrada
    add_p(ref_p, "Supuestos sobre los datos de entrada y qué ocurriría si no se cumplen:", bold=True, size_pt=10.0, space_before=6, space_after=2)
    add_p(ref_p, "La ventaja asintótica O(n) de los algoritmos no comparativos depende de premisas estrictas sobre los datos. Si estas premisas se violan, su ventaja se desvanece:")
    
    add_custom_p(ref_p, [
        ("• Supuesto de Radix Sort: ", True, False, 10.0, None),
        ("Asume que la longitud de las claves d es fija, pequeña e independiente de n (en el taller, d = 8 dígitos).", False, False, 10.0, None)
    ])
    add_custom_p(ref_p, [
        ("  - ¿Qué ocurre si son números de 16 dígitos?: ", False, True, 10.0, None),
        ("Dado que la complejidad es O(d · (n + b)), duplicar la cantidad de dígitos de 8 a 16 obliga al algoritmo a realizar exactamente el doble de pasadas completas de Counting Sort. Por ende, el tiempo de ejecución aumentaría prácticamente al doble (~2×).", False, False, 10.0, None)
    ])
    add_custom_p(ref_p, [
        ("  - ¿Qué ocurre si la longitud de la clave escala con n?: ", False, True, 10.0, None),
        ("Si las claves son enteros arbitrarios de tamaño n^c o cadenas de texto largas donde d = Θ(log n), la complejidad pasa a ser O(n log n). En este escenario, Radix Sort pierde su ventaja teórica lineal y resulta a menudo más lento que QuickSort o Merge Sort debido al costo constante de gestionar arreglos auxiliares y operaciones de copia.", False, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("• Supuesto de Bucket Sort: ", True, False, 10.0, None),
        ("Asume que los datos están distribuidos de forma homogénea y uniforme a lo largo del intervalo [min, max]. En el taller, esto se cumplió perfectamente con datos generados uniformemente entre 10.000.000 y 99.999.999 repartidos en 100.000 cubetas.", False, False, 10.0, None)
    ])
    add_custom_p(ref_p, [
        ("  - ¿Qué ocurre si los valores se concentran en un rango pequeño (distribución sesgada)?: ", False, True, 10.0, None),
        ("Si los datos presentan una distribución asimétrica, normal con baja varianza o valores masivamente repetidos, la función de dispersión colapsa: la inmensa mayoría de los elementos caerá en una sola cubeta o en un puñado de ellas. Como el algoritmo interno que ordena cada cubeta es Insertion Sort (cuadrático), si una cubeta recibe m ≈ n elementos, el tiempo de ordenamiento interno se dispara a O(m²) ≈ O(n²).", False, False, 10.0, None)
    ])
    add_custom_p(ref_p, [
        ("  - Impacto comparativo: ", False, True, 10.0, None),
        ("En dicho peor caso, Bucket Sort pasaría de 0,5 segundos a más de 78 segundos en Java, y de 1,9 segundos a más de 8 horas en Python, rindiendo drásticamente peor que algoritmos comparativos robustos como Merge Sort o Heap Sort, cuya cota O(n log n) es estrictamente inmune a la distribución de los datos.", False, False, 10.0, None)
    ])
    
    # --- SUBSECCIÓN 2 ---
    add_p(ref_p, "2. ¿Ordenar o no ordenar? (Punto de Equilibrio k*)", bold=True, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=4, heading_style="Heading 1")
    
    add_p(ref_p, "Supóngase que se dispone de un arreglo desordenado de 1.000.000 de elementos sobre el cual se deben ejecutar k búsquedas. Se analiza formalmente a partir de qué valor de k resulta más conveniente ordenar el arreglo previamente con el algoritmo más eficiente medido y utilizar Búsqueda Binaria, en lugar de realizar k búsquedas lineales sobre el arreglo en desorden.")
    
    # Formulación matemática
    add_p(ref_p, "Planteamiento formal del modelo de costo amortizado:", bold=True, size_pt=10.0, space_before=4, space_after=2)
    
    add_custom_p(ref_p, [
        ("• Estrategia A (Búsqueda Lineal sin ordenar): ", True, False, 10.0, None),
        ("No se invierte ningún tiempo previo en preparar los datos (T_prep = 0). Cada una de las k consultas de búsqueda se realiza secuencialmente sobre el arreglo desordenado. Si t_lineal es el tiempo promedio de una búsqueda lineal:\n", False, False, 10.0, None),
        ("T_A(k) = k · t_lineal", True, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("• Estrategia B (Ordenar previamente y aplicar Búsqueda Binaria): ", True, False, 10.0, None),
        ("Se invierte un tiempo inicial único de ordenamiento T_ordenar sobre el millón de elementos. Una vez ordenado, cada una de las k búsquedas se realiza en tiempo logarítmico con Búsqueda Binaria (t_binaria):\n", False, False, 10.0, None),
        ("T_B(k) = T_ordenar + k · t_binaria", True, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("• Condición de superioridad (Break-even): ", True, False, 10.0, None),
        ("Ordenar resulta más conveniente cuando el tiempo total acumulado de la Estrategia B es estrictamente menor al de la Estrategia A:\n", False, False, 10.0, None),
        ("T_B(k) < T_A(k)  ⟺  T_ordenar + k · t_binaria < k · t_lineal\n", False, True, 10.0, None),
        ("k · (t_lineal - t_binaria) > T_ordenar  ⟺  k > T_ordenar / (t_lineal - t_binaria)", True, False, 10.0, None)
    ])
    
    add_p(ref_p, "El punto de corte exacto es k* = T_ordenar / (t_lineal - t_binaria). Dado que el número de búsquedas k debe ser un valor entero positivo, la condición mínima requerida es k ≥ ⌈k*⌉ (o ⌊k*⌋ + 1).")
    
    # Tiempos base
    add_p(ref_p, "Tabla 18. Tiempos base medidos para n = 1.000.000 empleados en el cálculo", bold=True, size_pt=9.0, space_before=4, space_after=2)
    
    tbl18 = doc.add_table(rows=5, cols=6)
    ref_p._p.addprevious(tbl18._tbl)
    
    headers18 = ["Lenguaje", "Algoritmo de Ordenamiento", "T_ordenar (ms)", "t_lineal (por búsqueda)", "t_binaria (por búsqueda)", "Δt ahorrado (t_lin - t_bin)"]
    filas18 = [
        ["Java", "Radix Sort (Mejor absoluto)", "83,320 ms", "149,00 µs (0,14900 ms)", "274,3 ns (0,0002743 ms)", "0,148726 ms"],
        ["Java", "Dual-Pivot QS (Mejor comparativo)", "91,854 ms", "149,00 µs (0,14900 ms)", "274,3 ns (0,0002743 ms)", "0,148726 ms"],
        ["Python", "Bucket Sort (Mejor absoluto)", "1.915,109 ms", "26,440 ms", "5,64 µs (0,00564 ms)", "26,43436 ms"],
        ["Python", "Dual-Pivot QS (Mejor comparativo)", "2.600,673 ms", "26,440 ms", "5,64 µs (0,00564 ms)", "26,43436 ms"]
    ]
    for c_idx, h_text in enumerate(headers18):
        tbl18.cell(0, c_idx).text = h_text
    for r_idx, fila in enumerate(filas18):
        for c_idx, val in enumerate(fila):
            tbl18.cell(r_idx + 1, c_idx).text = val
    col_w18 = [Inches(0.9), Inches(1.8), Inches(1.1), Inches(1.2), Inches(1.2), Inches(1.2)]
    col_a18 = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT]
    style_table(tbl18, col_w18, col_a18)
    
    # Cálculos numéricos detallados
    add_p(ref_p, "Desarrollo del cálculo paso a paso:", bold=True, size_pt=10.0, space_before=6, space_after=2)
    
    add_custom_p(ref_p, [
        ("1. En Java con el mejor algoritmo absoluto (Radix Sort):\n", True, False, 10.0, None),
        ("k* = 83,320 ms / (0,14900 ms - 0,0002743 ms) = 83,320 / 0,1487257 ≈ 560,23\n", False, False, 10.0, None),
        ("Resultado: A partir de k ≥ 561 búsquedas, ordenar con Radix Sort y usar Búsqueda Binaria es más rápido que la búsqueda lineal.", True, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("2. En Java con el mejor algoritmo comparativo (Dual-Pivot QuickSort):\n", True, False, 10.0, None),
        ("k* = 91,854 ms / 0,1487257 ms ≈ 617,61\n", False, False, 10.0, None),
        ("Resultado: A partir de k ≥ 618 búsquedas, ordenar con Dual-Pivot QuickSort y usar Búsqueda Binaria supera a la búsqueda lineal.", True, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("3. En Python con el mejor algoritmo absoluto (Bucket Sort):\n", True, False, 10.0, None),
        ("k* = 1.915,109 ms / (26,440 ms - 0,00564 ms) = 1.915,109 / 26,43436 ≈ 72,45\n", False, False, 10.0, None),
        ("Resultado: A partir de k ≥ 73 búsquedas, ordenar con Bucket Sort y usar Búsqueda Binaria es superior a la búsqueda lineal.", True, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("4. En Python con el mejor algoritmo comparativo (Dual-Pivot QuickSort):\n", True, False, 10.0, None),
        ("k* = 2.600,673 ms / 26,43436 ms ≈ 98,38\n", False, False, 10.0, None),
        ("Resultado: A partir de k ≥ 99 búsquedas, ordenar con Dual-Pivot QuickSort y usar Búsqueda Binaria supera a la búsqueda lineal.", True, False, 10.0, None)
    ])
    
    # Tabla 19
    add_p(ref_p, "Tabla 19. Punto de equilibrio k* (búsquedas mínimas para justificar el ordenamiento)", bold=True, size_pt=9.0, space_before=4, space_after=2)
    tbl19 = doc.add_table(rows=5, cols=6)
    ref_p._p.addprevious(tbl19._tbl)
    
    headers19 = ["Lenguaje", "Algoritmo de Ordenamiento", "T_ordenar", "k* Exacto", "k Mínimo Entero", "Fracción de n (1.000.000)"]
    filas19 = [
        ["Java", "Radix Sort (Mejor absoluto)", "83,32 ms", "560,23", "k ≥ 561 búsquedas", "0,056% de los datos"],
        ["Java", "Dual-Pivot QS (Mejor comparativo)", "91,85 ms", "617,61", "k ≥ 618 búsquedas", "0,062% de los datos"],
        ["Python", "Bucket Sort (Mejor absoluto)", "1.915,11 ms", "72,45", "k ≥ 73 búsquedas", "0,007% de los datos"],
        ["Python", "Dual-Pivot QS (Mejor comparativo)", "2.600,67 ms", "98,38", "k ≥ 99 búsquedas", "0,010% de los datos"]
    ]
    for c_idx, h_text in enumerate(headers19):
        tbl19.cell(0, c_idx).text = h_text
    for r_idx, fila in enumerate(filas19):
        for c_idx, val in enumerate(fila):
            tbl19.cell(r_idx + 1, c_idx).text = val
    col_w19 = [Inches(1.0), Inches(2.0), Inches(1.1), Inches(1.0), Inches(1.4), Inches(1.2)]
    col_a19 = [WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.LEFT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.RIGHT, WD_ALIGN_PARAGRAPH.CENTER, WD_ALIGN_PARAGRAPH.RIGHT]
    style_table(tbl19, col_w19, col_a19)
    
    # Figura 11
    p_img11 = ref_p.insert_paragraph_before()
    p_img11.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_img11.paragraph_format.space_before = Pt(6)
    p_img11.paragraph_format.space_after = Pt(2)
    run11 = p_img11.add_run()
    run11.add_picture(GRAFICA2_PATH, width=Inches(6.3))
    
    add_p(ref_p, "Figura 11. Punto de equilibrio (Break-even): Costo acumulado T(k) vs número de búsquedas k en Java y Python.", italic=True, size_pt=9.0, align=WD_ALIGN_PARAGRAPH.CENTER, space_before=2, space_after=6)
    
    # Análisis del resultado de k*
    add_p(ref_p, "Análisis de las implicaciones del punto de equilibrio:", bold=True, size_pt=10.0, space_before=6, space_after=2)
    
    add_custom_p(ref_p, [
        ("1. Umbral asombrosamente bajo: ", True, False, 10.0, None),
        ("Tanto en Java (k ≥ 561) como en Python (k ≥ 73), el punto de equilibrio ocurre en una fracción ínfima del tamaño total del arreglo (menos del 0,06% en Java y menos del 0,01% en Python). Esto significa que en prácticamente cualquier escenario de ingeniería o sistema de software del mundo real donde se realicen más de unas pocas centenas de consultas, no ordenar el arreglo es una decisión de diseño ineficiente.", False, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("2. ¿Por qué Python amortiza mucho antes que Java (73 frente a 561 búsquedas)?: ", True, False, 10.0, None),
        ("A primera vista parece contraintuitivo, pues ordenar en Python toma 1.915 ms frente a sólo 83,3 ms en Java. Sin embargo, la razón se encuentra en el denominador (la brecha de la búsqueda lineal): la búsqueda lineal en Python es abismalmente lenta (26,44 ms por búsqueda frente a 0,149 ms en Java, una penalización de 177×). En Python, cada búsqueda lineal adicional añade un retardo perceptible de más de 26 milisegundos, por lo que el costo del ordenamiento se paga en apenas 73 consultas.", False, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("3. ¿Qué ocurriría si se utiliza Búsqueda por Interpolación en lugar de Binaria?: ", True, False, 10.0, None),
        ("En el Punto 2, la Búsqueda por Interpolación resultó ser la más veloz de todas (33,8 ns en Java y 2,24 µs en Python). Al recalcular k* con este algoritmo: en Java k* = 83,320 / (0,14900 - 0,0000338) ≈ 559,32 (k ≥ 560), y en Python k* = 1.915,109 / (26,440 - 0,00224) ≈ 72,44 (k ≥ 73). La diferencia es prácticamente nula (a lo sumo 1 búsqueda menos), lo que demuestra que cualquier algoritmo logarítmico o sublogarítmico es tan extremadamente veloz frente a la búsqueda lineal que el tiempo de búsqueda sobre el arreglo ordenado resulta despreciable frente a la lentitud de la búsqueda secuencial.", False, False, 10.0, None)
    ])
    
    # --- SUBSECCIÓN 3 ---
    add_p(ref_p, "3. Conclusiones del Punto 3", bold=True, size_pt=10.5, align=WD_ALIGN_PARAGRAPH.LEFT, space_before=12, space_after=4, heading_style="Heading 1")
    
    add_custom_p(ref_p, [
        ("Superioridad de los algoritmos no comparativos bajo condiciones controladas: ", True, False, 10.0, None),
        ("Radix Sort en Java (83,3 ms) y Bucket Sort en Python (1.915 ms) ratificaron en el laboratorio su capacidad teórica de romper la cota inferior Ω(n log n). Sin embargo, son algoritmos especializados: dependen rígidamente de que la longitud de las claves sea constante o de que la distribución sea uniforme. Ante desviaciones de estos supuestos, colapsan a rendimientos subóptimos o cuadráticos O(n²), donde algoritmos comparativos como Dual-Pivot QuickSort o Merge Sort ofrecen una robustez inalcanzable.", False, False, 10.0, None)
    ])
    
    add_custom_p(ref_p, [
        ("Regla de decisión empírica para ordenar vs buscar en desorden: ", True, False, 10.0, None),
        ("La decisión de preprocesar datos mediante ordenamiento tiene una rentabilidad de inversión computacional inmediata: realizar apenas entre 73 (Python) y 561 (Java) búsquedas justifica por completo invertir en el ordenamiento del millón de datos. A partir de dicho punto, cada consulta subsiguiente se acelera en factores de entre 500× (Java) y 4.600× (Python), garantizando la escalabilidad del sistema.", False, False, 10.0, None)
    ])
    
    # -------------------------------------------------------------
    # SECCIÓN: USO DE INTELIGENCIA ARTIFICIAL (PUNTO 3)
    # -------------------------------------------------------------
    # Append at the very end of the document using doc.add_paragraph()
    p_ia = doc.add_paragraph()
    p_ia.alignment = WD_ALIGN_PARAGRAPH.LEFT
    p_ia.paragraph_format.space_before = Pt(8)
    p_ia.paragraph_format.space_after = Pt(3)
    p_ia.paragraph_format.line_spacing = 1.15
    run_ia = p_ia.add_run("Punto 3. Usamos Antigravity (Google) como apoyo en las siguientes partes:")
    run_ia.font.name = "Calibri"
    run_ia.font.size = Pt(10.0)
    run_ia.font.bold = True
    
    def append_doc_p(runs_data, space_before=2, space_after=2):
        p = doc.add_paragraph()
        p.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
        p.paragraph_format.space_before = Pt(space_before)
        p.paragraph_format.space_after = Pt(space_after)
        p.paragraph_format.line_spacing = 1.15
        for r_text, r_bold, r_italic, r_size, r_color in runs_data:
            r = p.add_run(r_text)
            r.font.name = "Calibri"
            r.font.size = Pt(r_size)
            r.font.bold = r_bold
            r.font.italic = r_italic
            if r_color:
                r.font.color.rgb = r_color
        return p
        
    append_doc_p([
        ("Implementación y cálculo: ", True, False, 10.0, None),
        ("Para diseñar los scripts en Python de análisis matemático, cómputo exacto de los puntos de equilibrio k* y generación de las visualizaciones comparativas de familias y curvas de amortización mediante matplotlib.", False, False, 10.0, None)
    ])
    
    append_doc_p([
        ("Redacción y formalización: ", True, False, 10.0, None),
        ("Para formular la demostración matemática rigurosa del límite inferior Ω(n log n) con el modelo de árboles de decisión y redactar la síntesis comparativa entre ordenamiento y búsqueda sustentada en las mediciones empíricas de los puntos 1 y 2.", False, False, 10.0, None)
    ])
    
    append_doc_p([
        ("Gráficas: ", True, False, 10.0, None),
        ("Para generar las visualizaciones de comparación de familias y análisis de break-even con matplotlib a partir de los resultados medidos.", False, False, 10.0, None)
    ])
    
    print("Guardando documento actualizado...")
    doc.save(DOC_PATH)
    print("¡Documento Word guardado exitosamente!")

if __name__ == "__main__":
    ejecutar()
